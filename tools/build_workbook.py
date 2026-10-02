"""Rebuild workbooks from TEST_INVENTORY.md + MCDC_MATRIX.md.
- workbook/Workbook.xlsx (working): Sheet1 8 cols, Sheet2 5 cols.
- workbook/Workbook_proposed.xlsx (submission candidate): Sheet1 8 cols,
  Sheet2 6 cols with col F 'Source / test reference'; gap rows labelled via col F.
"""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = "/home/zexr/Desktop/sqe_A2"


def md_tables(path):
    tables, section, cur = [], "", None
    with open(path) as f:
        for raw in f:
            line = raw.rstrip("\n")
            m = re.match(r"###\s+(.*)", line)
            if m:
                if cur:
                    tables.append((section, cur))
                section = m.group(1).strip()
                cur = None
            elif line.startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?", c or "") for c in cells):
                    continue
                if cur is None:
                    cur = {"header": cells, "rows": []}
                else:
                    cur["rows"].append(cells)
    if cur:
        tables.append((section, cur))
    return tables


# ---- Sheet 1 data ----
inv_tables = md_tables(f"{BASE}/TEST_INVENTORY.md")
inv = []
for _, t in inv_tables:
    if t["header"] and t["header"][0] == "Test ID":
        inv += [r for r in t["rows"] if len(r) == 8 and re.fullmatch(r"(HST|BST|PST)-\d+", r[0])]
assert len(inv) == 61, f"expected 61 inventory rows, got {len(inv)}"

# ---- Sheet 2 data ----
mx_tables = md_tables(f"{BASE}/MCDC_MATRIX.md")


def src_ref(section):
    m = re.search(r"(\w+\.cpp)\s+L(\d+)", section)
    if m:
        return f"{m.group(1)}:{m.group(2)}"
    m = re.search(r"(\w+\.cpp)\s+L(\d+).*L(\d+)", section)
    if m:
        return f"{m.group(1)}:{m.group(2)}"
    m = re.search(r"(\w+\.cpp)", section)
    return m.group(1) if m else section.split("—")[0].strip()


mx = []  # (decision, conds, outcome, pair, demonstrates, srcref)
for section, t in mx_tables:
    h = t["header"]
    try:
        oi = next(i for i, c in enumerate(h) if "outcome" in c.lower())
    except StopIteration:
        continue
    dec = section.split("—")[0].strip()
    ref = src_ref(section)
    for r in t["rows"]:
        if len(r) != len(h):
            continue
        conds = "; ".join(f"{h[i]}={r[i]}" for i in range(oi))
        pair = r[oi + 1] if oi + 1 < len(r) else ""
        dem = r[oi + 2] if oi + 2 < len(r) else ""
        if pair.strip() in ("—", ""):
            colf = f"GAP (investigated, not a pair): {ref} — see report §7"
        else:
            colf = f"{ref} / {pair}"
        mx.append((dec, conds, r[oi], pair, dem, colf))
assert len(mx) == 32, f"expected 32 mcdc rows, got {len(mx)}"


def style_sheet(ws, widths):
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Border(*[Side(style="thin", color="B0B0B0")] * 4)
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = wrap
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
    ws.freeze_panes = "A2"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def build(path, six_col):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Test Inventory"
    ws1.append(["Test ID", "Component / function", "Purpose / scenario",
                "Key controlled input / state", "Expected result", "Exec",
                "Structural-coverage target", "Test-file / evidence reference"])
    for r in inv:
        ws1.append(r)
    style_sheet(ws1, [12, 24, 30, 32, 30, 10, 26, 44])
    ws2 = wb.create_sheet("MC_DC Evidence")
    if six_col:
        ws2.append(["Compound decision (file:line)", "Atomic-condition values",
                    "Overall decision outcome", "Independence pair (tests)",
                    "Condition demonstrated independent", "Source / test reference"])
        for r in mx:
            ws2.append(list(r))
        style_sheet(ws2, [30, 38, 24, 28, 36, 34])
    else:
        ws2.append(["Compound decision (file:line)", "Atomic-condition values",
                    "Overall decision outcome", "Independence pair (tests)",
                    "Condition demonstrated independent"])
        for r in mx:
            ws2.append(list(r[:5]))
        style_sheet(ws2, [34, 40, 26, 30, 40])
    wb.save(path)
    print(f"saved {path}: sheet1={ws1.max_row-1}, sheet2={ws2.max_row-1}")


build(f"{BASE}/workbook/Workbook.xlsx", six_col=False)
build(f"{BASE}/workbook/Workbook_proposed.xlsx", six_col=True)
